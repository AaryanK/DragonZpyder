"""Regression tests for attaching images and documents to the existing Mesh chat.

No OpenRouter requests are made; authentication and billing remain in /mesh routes.
"""
import base64
import io
import sys
from pathlib import Path
from docx import Document
import pytest
from fastapi import HTTPException

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from backend.mesh import document_part, messages_from, MAX_BODY


def uri(mime: str, data: bytes) -> str:
    return "data:" + mime + ";base64," + base64.b64encode(data).decode("ascii")


def attachment(name: str, mime: str, data: bytes):
    return {"type": "file", "file": {"filename": name, "file_data": uri(mime, data)}}


def test_pdf_is_forwarded_with_openrouter_file_schema():
    file = attachment("guide.pdf", "application/pdf", b"%PDF-1.4\nsample")
    item, size, pdf = document_part(file)
    assert pdf and size == len(b"%PDF-1.4\nsample")
    assert item["file"]["filename"] == "guide.pdf"
    messages, vision, has_pdf = messages_from({"messages": [
        {"role": "user", "content": [{"type": "text", "text": "Summarize"}, file]}
    ]})
    assert not vision and has_pdf
    assert messages[0]["content"][1]["type"] == "file"


def test_docx_is_extracted_into_bounded_text():
    document = Document()
    document.add_paragraph("Quarterly plan")
    document.add_paragraph("Finish testing")
    data = io.BytesIO()
    document.save(data)
    part, _, pdf = document_part(attachment(
        "plan.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", data.getvalue()
    ))
    assert not pdf
    assert "Quarterly plan" in part["text"]
    assert "Finish testing" in part["text"]


def test_plain_text_and_windows_filename():
    part, size, pdf = document_part(attachment("C:\\fakepath\\notes.txt", "text/plain", b"line one"))
    assert part == {"type": "text", "text": "Attached file (notes.txt):\nline one"}
    assert size == 8 and not pdf


def test_image_and_document_in_single_chat_request():
    image = {"type": "image_url", "image_url": {"url": uri("image/png", b"\x89PNG\r\n\x1a\nimage")}}
    txt = attachment("notes.md", "text/plain", b"read these notes")
    body = {"messages": [{"role": "user", "content": [
        {"type": "text", "text": "Explain these"}, image, txt
    ]}]}
    messages, vision, has_pdf = messages_from(body)
    assert vision and not has_pdf
    assert messages[0]["content"][1]["type"] == "image_url"
    assert "read these notes" in messages[0]["content"][2]["text"]


def test_rejects_wrong_file_formats_and_unbounded_uploads():
    with pytest.raises(HTTPException) as wrong:
        document_part(attachment("fake.pdf", "application/pdf", b"Not a PDF"))
    assert wrong.value.status_code == 400
    with pytest.raises(HTTPException) as too_many:
        messages_from({"messages": [{"role": "user", "content": [
            {"type": "text", "text": "test"},
            *[attachment(f"{n}.txt", "text/plain", b"data") for n in range(4)]
        ]}]})
    assert too_many.value.status_code == 413
    assert MAX_BODY <= 12 * 1024 * 1024


def test_assistant_cannot_supply_attachment():
    with pytest.raises(HTTPException) as invalid:
        messages_from({"messages": [{"role": "assistant", "content": [
            attachment("tool.txt", "text/plain", b"forged")
        ]}]})
    assert invalid.value.status_code == 400
