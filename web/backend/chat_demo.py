"""Password-gated multimodal Chat Lab using DragonZpyder central inference."""
from __future__ import annotations

import base64
import binascii
import hmac
import hashlib
import json
import os
import re
import time
from collections import defaultdict, deque
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit
from zipfile import BadZipFile, ZipFile

import httpx
from docx import Document
from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel, ConfigDict, Field, ValidationError

router = APIRouter()
ROOT = Path(__file__).parent / "static"
COOKIE = "dz_demo_session"
GATEWAY = os.getenv("DRAGONZPYDER_API_ORIGIN", "http://dragonzpyder-api.railway.internal:8080").rstrip("/") + "/internal/provider/chat/completions"
MAX_BODY = 9 * 1024 * 1024
MAX_SINGLE = 3 * 1024 * 1024
MAX_TEXT = 65_000
DEFAULT_MODEL = "google/gemini-2.5-flash"
MODEL_PATTERN = re.compile(r"^[a-zA-Z0-9_.:/-]{3,110}$")
REQUEST_TIMES: dict[str, deque[float]] = defaultdict(deque)
LOGIN_FAILURES: deque[float] = deque(maxlen=40)


class ChatTurn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    role: str
    content: str = Field(min_length=1, max_length=10_000)


class Attachment(BaseModel):
    model_config = ConfigDict(extra="forbid")
    name: str = Field(min_length=1, max_length=140)
    data: str = Field(min_length=4, max_length=4_300_000)


class ChatBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    model: str = Field(min_length=3, max_length=110)
    messages: list[ChatTurn] = Field(min_length=1, max_length=12)
    attachments: list[Attachment] = Field(default_factory=list, max_length=3)


def models() -> list[str]:
    raw = os.getenv("DRAGONZPYDER_CHAT_MODELS", DEFAULT_MODEL)
    selected = list(dict.fromkeys(s.strip() for s in raw.split(",") if MODEL_PATTERN.fullmatch(s.strip())))
    return selected[:8] or [DEFAULT_MODEL]


def configured() -> bool:
    return len(os.getenv("DRAGONZPYDER_CHAT_GATEWAY_TOKEN", "")) >= 32 and len(os.getenv("DRAGONZPYDER_CHAT_PASSWORD", "")) >= 16


def _signature(expiry: str) -> str:
    secret = os.getenv("DRAGONZPYDER_CHAT_PASSWORD", "")
    return hmac.new(secret.encode(), ("dragonzpyder-demo-v1:" + expiry).encode(), hashlib.sha256).hexdigest()


def _authenticated(request: Request) -> bool:
    if not configured():
        return False
    value = request.cookies.get(COOKIE, "")
    try:
        expiry, signature = value.split(".", 1)
        return (expiry.isascii() and expiry.isdigit() and len(expiry) <= 12
                and int(expiry) > int(time.time())
                and hmac.compare_digest(_signature(expiry), signature))
    except (ValueError, TypeError):
        return False


def _protect_mutation(request: Request) -> None:
    if request.headers.get("x-dragonzpyder-chat") != "1":
        raise HTTPException(403, "Demo client header required")
    if request.headers.get("sec-fetch-site", "same-origin") not in ("same-origin", "none"):
        raise HTTPException(403, "Cross-site requests are not allowed")
    origin = request.headers.get("origin")
    if origin:
        parsed = urlsplit(origin)
        if parsed.scheme not in ("https", "http") or parsed.netloc.lower() != request.headers.get("host", "").lower():
            raise HTTPException(403, "Cross-origin requests are not allowed")


def _require_auth(request: Request) -> str:
    if not _authenticated(request):
        raise HTTPException(401, "Sign in to use the chat demo")
    # All signed sessions share this demo quota. This is deliberately not multi-tenant billing.
    return "demo"


def _rate_limit(bucket: str, max_calls: int = 12) -> None:
    now = time.monotonic()
    q = REQUEST_TIMES[bucket]
    while q and q[0] <= now - 3600:
        q.popleft()
    if len(q) >= max_calls:
        raise HTTPException(429, "Demo hourly limit reached")
    q.append(now)


@router.get("/chat")
async def page() -> FileResponse:
    return FileResponse(ROOT / "index.html", media_type="text/html", headers={"Cache-Control": "no-store"})


@router.get("/chat/app.css")
async def style() -> FileResponse:
    return FileResponse(ROOT / "app.css", media_type="text/css", headers={"Cache-Control": "no-store"})


@router.get("/chat/app.js")
async def script() -> FileResponse:
    return FileResponse(ROOT / "app.js", media_type="text/javascript", headers={"Cache-Control": "no-store"})


@router.get("/demo/session")
async def session(request: Request) -> dict:
    return {"authenticated": _authenticated(request), "configured": configured(),
            "models": models() if _authenticated(request) else []}


@router.post("/demo/login")
async def login(request: Request) -> JSONResponse:
    _protect_mutation(request)
    if not configured():
        raise HTTPException(503, "Set DRAGONZPYDER_CHAT_GATEWAY_TOKEN and DRAGONZPYDER_CHAT_PASSWORD (16+ characters) in Railway")
    raw = await request.body()
    if len(raw) > 512:
        raise HTTPException(413, "Login request too large")
    try:
        supplied = json.loads(raw)["password"]
    except (ValueError, KeyError, TypeError):
        raise HTTPException(400, "Password required")
    if not isinstance(supplied, str):
        raise HTTPException(400, "Password required")
    now = time.monotonic()
    while LOGIN_FAILURES and LOGIN_FAILURES[0] < now - 600:
        LOGIN_FAILURES.popleft()
    if len(LOGIN_FAILURES) >= 20:
        raise HTTPException(429, "Too many login attempts; try later")
    expected = os.environ["DRAGONZPYDER_CHAT_PASSWORD"]
    if not hmac.compare_digest(supplied.encode(), expected.encode()):
        LOGIN_FAILURES.append(now)
        raise HTTPException(401, "Incorrect demo password")
    expiry = str(int(time.time()) + 8 * 3600)
    res = JSONResponse({"authenticated": True, "models": models()})
    res.set_cookie(COOKIE, expiry + "." + _signature(expiry), max_age=8 * 3600,
                   httponly=True, secure=True, samesite="strict", path="/")
    return res


@router.post("/demo/logout")
async def logout(request: Request) -> JSONResponse:
    _protect_mutation(request)
    res = JSONResponse({"authenticated": False})
    res.delete_cookie(COOKIE, path="/")
    return res


def _attachment_parts(items: list[Attachment]) -> tuple[list[dict], bool]:
    parts: list[dict] = []
    pdf_used = False
    total = 0
    for item in items:
        filename = item.name.replace("\\", "/").rsplit("/", 1)[-1]
        suffix = Path(filename).suffix.lower()
        try:
            data = base64.b64decode(item.data, validate=True)
        except (binascii.Error, ValueError):
            raise HTTPException(422, "Invalid attachment encoding")
        total += len(data)
        if not data or len(data) > MAX_SINGLE or total > 7 * 1024 * 1024:
            raise HTTPException(413, "Each attachment must be at most 3 MB (7 MB combined)")
        if suffix in (".png", ".jpg", ".jpeg", ".webp"):
            if suffix == ".png" and not data.startswith(b"\x89PNG\r\n\x1a\n"):
                raise HTTPException(422, "Invalid PNG image")
            if suffix in (".jpg", ".jpeg") and not data.startswith(b"\xff\xd8\xff"):
                raise HTTPException(422, "Invalid JPEG image")
            if suffix == ".webp" and not (data.startswith(b"RIFF") and data[8:12] == b"WEBP"):
                raise HTTPException(422, "Invalid WebP image")
            media = {".png":"image/png", ".jpg":"image/jpeg", ".jpeg":"image/jpeg", ".webp":"image/webp"}[suffix]
            parts.append({"type":"image_url", "image_url":{"url":f"data:{media};base64,{item.data}"}})
        elif suffix == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise HTTPException(422, "Invalid PDF file")
            pdf_used = True
            parts.append({"type":"file", "file":{"filename":filename,
                                                   "file_data":"data:application/pdf;base64," + item.data}})
        elif suffix == ".docx":
            if len(data) > 1_000_000:
                raise HTTPException(413, "Word documents must be under 1 MB")
            try:
                with ZipFile(BytesIO(data)) as archive:
                    if sum(x.file_size for x in archive.infolist()) > 5_000_000:
                        raise HTTPException(413, "Word document expands too large")
                doc = Document(BytesIO(data))
                txt = "\n".join(p.text for p in doc.paragraphs)
                for table in doc.tables:
                    for row in table.rows:
                        txt += "\n" + " | ".join(cell.text for cell in row.cells)
            except (BadZipFile, ValueError, KeyError, OSError):
                raise HTTPException(422, "Invalid Word document")
            parts.append({"type":"text", "text":f"Attached Word document ({filename}):\n{txt[:MAX_TEXT]}"})
        elif suffix in (".txt", ".md", ".csv", ".json", ".py", ".log"):
            if len(data) > 250_000:
                raise HTTPException(413, "Text attachments must be under 250 KB")
            try:
                txt = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise HTTPException(422, "Text file must use UTF-8")
            parts.append({"type":"text", "text":f"Attached text file ({filename}):\n{txt[:MAX_TEXT]}"})
        else:
            raise HTTPException(415, "Supported uploads: PNG, JPEG, WebP, PDF, DOCX, TXT, MD, CSV, JSON, PY, LOG")
    return parts, pdf_used


def _payload(body: ChatBody) -> dict:
    allowed = models()
    if body.model not in allowed:
        raise HTTPException(403, "Model is not enabled for this demo")
    if body.messages[-1].role != "user" or any(t.role not in ("user", "assistant") for t in body.messages):
        raise HTTPException(422, "Messages must contain user and assistant turns, ending with a user turn")
    parts, has_pdf = _attachment_parts(body.attachments)
    messages = [{"role":t.role, "content":t.content} for t in body.messages[:-1]]
    final = body.messages[-1].content
    if parts:
        messages.append({"role":"user", "content":[{"type":"text","text":final}, *parts]})
    else:
        messages.append({"role":"user", "content":final})
    payload = {"model": body.model, "messages":messages, "max_tokens":1024,
               "stream":False}
    if has_pdf:
        # Avoid the more expensive default OCR engine in this small prototype.
        payload["plugins"] = [{"id":"file-parser", "pdf":{"engine":"cloudflare-ai"}}]
    return payload


@router.post("/demo/chat")
async def chat(request: Request) -> dict:
    _protect_mutation(request)
    principal = _require_auth(request)
    # Read with an explicit body limit before JSON parsing, so base64 attachments cannot balloon memory.
    data = bytearray()
    async for chunk in request.stream():
        data.extend(chunk)
        if len(data) > MAX_BODY:
            raise HTTPException(413, "Request exceeds 9 MB")
    try:
        body = ChatBody.model_validate_json(bytes(data))
    except ValidationError:
        raise HTTPException(422, "Invalid chat request")
    payload = _payload(body)
    _rate_limit(principal)
    headers = {"Authorization": "Bearer " + os.environ["DRAGONZPYDER_CHAT_GATEWAY_TOKEN"],
               "Content-Type": "application/json"}
    try:
        upstream = await request.app.state.http.post(GATEWAY, json=payload, headers=headers, timeout=90)
    except httpx.TimeoutException:
        raise HTTPException(504, "DragonZpyder inference timed out")
    except httpx.RequestError:
        raise HTTPException(502, "Could not reach DragonZpyder inference gateway")
    if upstream.status_code != 200:
        code = 429 if upstream.status_code == 429 else 502
        raise HTTPException(code, "DragonZpyder inference rejected the request (verify model, balance or limits)")
    try:
        item = upstream.json()
        answer = item["choices"][0]["message"]["content"]
        if isinstance(answer, list):
            answer = "\n".join(x.get("text", "") for x in answer if isinstance(x, dict) and x.get("type") == "text")
        if not isinstance(answer, str):
            raise ValueError("Non-text output")
        usage = item.get("usage") or {}
        safe_usage = {k:int(v) for k,v in usage.items() if k in ("prompt_tokens","completion_tokens","total_tokens") and type(v) is int}
        return {"reply":answer[:50_000], "model":str(item.get("model") or body.model)[:120], "usage":safe_usage}
    except (ValueError, KeyError, IndexError, TypeError):
        raise HTTPException(502, "DragonZpyder inference returned an unsupported response")
